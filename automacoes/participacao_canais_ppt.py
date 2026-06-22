import base64
import json
import os
import re
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt


SLIDE_IMAGE_BOXES = {
    "receita": (250000, 1400000, 11692000, 5150000),
    "pax": (1407685, 1111465, 9029680, 5109854),
    "tm": (1071776, 1194697, 10191025, 4713697),
}
UPDATE_TEXT_FONT_PT = 12

MONTHS_PT = {
    1: "Janeiro",
    2: "Fevereiro",
    3: "Março",
    4: "Abril",
    5: "Maio",
    6: "Junho",
    7: "Julho",
    8: "Agosto",
    9: "Setembro",
    10: "Outubro",
    11: "Novembro",
    12: "Dezembro",
}

MONTH_TOKENS_PT = {
    1: ["JANEIRO", "JAN"],
    2: ["FEVEREIRO", "FEV"],
    3: ["MARCO", "MAR"],
    4: ["ABRIL", "ABR"],
    5: ["MAIO", "MAI"],
    6: ["JUNHO", "JUN"],
    7: ["JULHO", "JUL"],
    8: ["AGOSTO", "AGO"],
    9: ["SETEMBRO", "SET"],
    10: ["OUTUBRO", "OUT"],
    11: ["NOVEMBRO", "NOV"],
    12: ["DEZEMBRO", "DEZ"],
}


def get_params():
    if len(sys.argv) <= 1:
        return {}
    return json.loads(base64.b64decode(sys.argv[1]).decode("utf-8"))


def normalize_update_date(value):
    text = str(value or "").strip()
    match = re.search(r"\b(\d{2}/\d{2}/\d{4})\b", text)
    if match:
        return match.group(1)
    return (datetime.now() - timedelta(days=1)).strftime("%d/%m/%Y")


def strip_accents(value):
    return "".join(
        ch for ch in unicodedata.normalize("NFD", str(value or ""))
        if unicodedata.category(ch) != "Mn"
    )


def resolve_report_period(month_label):
    now = datetime.now()
    text = str(month_label or "").strip()
    normalized = strip_accents(text).upper()

    month = now.month
    year = now.year

    for candidate_month, tokens in MONTH_TOKENS_PT.items():
        if any(token in normalized for token in tokens):
            month = candidate_month
            break

    year_match = re.search(r"\b(20\d{2})\b", text)
    if year_match:
        year = int(year_match.group(1))
    else:
        short_year_match = re.search(r"(?:^|[^\d])(\d{2})(?:$|[^\d])", text)
        if short_year_match:
            year = 2000 + int(short_year_match.group(1))

    return MONTHS_PT.get(month, MONTHS_PT[now.month]), year


def replace_text_preserving_runs(shape, replacements, update_date):
    if not getattr(shape, "has_text_frame", False):
        return

    full_text = shape.text_frame.text or ""
    if "Atualizado até dia" in full_text or "Atualizado ate dia" in full_text:
        shape.text_frame.text = f"Atualizado até dia {update_date}"
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = Pt(UPDATE_TEXT_FONT_PT)
            for run in paragraph.runs:
                run.font.size = Pt(UPDATE_TEXT_FONT_PT)
        return

    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            text = run.text
            for old, new in replacements.items():
                text = text.replace(old, new)
            run.text = text


def fit_picture(slide, image_path, box):
    x, y, w, h = box
    pic = slide.shapes.add_picture(str(image_path), x, y)
    original_w = pic.width
    original_h = pic.height
    scale = min(w / original_w, h / original_h)
    pic.width = int(original_w * scale)
    pic.height = int(original_h * scale)
    pic.left = int(x + (w - pic.width) / 2)
    pic.top = int(y + (h - pic.height) / 2)
    return pic


def main():
    params = get_params()
    template_path = Path(params["template_path"])
    output_path = Path(params["output_path"])
    images = params.get("images") or {}

    mes, report_year = resolve_report_period(params.get("month_label"))
    ano = str(report_year)
    ano_anterior = str(report_year - 1)
    mes = mes.capitalize()
    update_date = normalize_update_date(params.get("update_info"))

    replacements = {
        "Mês": mes,
        "Mes": mes,
        "Ano_anterior": ano_anterior,
        "Ano": ano,
    }

    prs = Presentation(str(template_path))

    for slide in prs.slides:
        for shape in slide.shapes:
            replace_text_preserving_runs(shape, replacements, update_date)

    if len(prs.slides) >= 4:
        slide_map = {
            "receita": prs.slides[1],
            "pax": prs.slides[2],
            "tm": prs.slides[3],
        }
        for key, slide in slide_map.items():
            image_path = images.get(key)
            if image_path and Path(image_path).exists():
                fit_picture(slide, Path(image_path), SLIDE_IMAGE_BOXES[key])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(json.dumps({
        "arquivo_principal": str(output_path),
        "arquivos_saida": [str(output_path)],
        "pasta_final": str(output_path.parent),
        "mensagem": "Apresentação de Participação de Canais gerada com sucesso."
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
